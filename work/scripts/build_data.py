"""Build site/data/data.json from the files in inputs/ (which are never edited).

Run from the project root:  python3 work/scripts/build_data.py
"""
import collections
import json
import re
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "inputs/data/Coding_Book latest.xlsx"
DOCX = ROOT / "inputs/data/For supp material_water paper (1).docx"
OUT = ROOT / "site/data/data.json"

# Expansions of the action-situation codes. The coding book does not spell
# these out; they follow common usage in the action-situation literature and
# need confirming with the author.
ACTION_SITUATIONS = [
    ("ADM", "Administration", "Who organises, approves, licenses and enforces."),
    ("PRV", "Provision", "Protecting water bodies and keeping infrastructure working."),
    ("APR", "Appropriation", "Who may take and use water, and on what terms."),
    ("MNT", "Monitoring", "Measuring, recording and reporting on water."),
    ("BP", "Basin planning", "Plans and councils at river-basin level."),
]

GROUPS = [
    ("central", "Central government"),
    ("agency", "Committees and national services"),
    ("basin", "Basin and local bodies"),
    ("users", "Water users, public and environment"),
]

# code -> (group, short label). Full names come from the codebook sheet.
ACTORS = {
    "Government_KZ": ("central", "Government"),
    "PrimeMinister_KZ": ("central", "Prime Minister"),
    "MVRI": ("central", "Water Ministry"),
    "MA": ("central", "Agriculture Ministry"),
    "MNE": ("central", "Economy Ministry"),
    "MIC": ("central", "Industry Ministry"),
    "MENR": ("central", "Ecology Ministry"),
    "ME": ("central", "Energy Ministry"),
    "MH": ("central", "Health Ministry"),
    "MinFin_KZ": ("central", "Finance Ministry"),
    "State": ("central", "State (general)"),
    "Water_Council": ("central", "Water Council"),
    "Subsoil_Authority": ("central", "Subsoil authority"),
    "Civil_Protection_Authority": ("central", "Civil protection"),
    "Housing_Utilities_Authority": ("central", "Housing and utilities"),
    "AgroIndustry_Authority": ("central", "Agro-industry authority"),
    "Forestry_Authority": ("central", "Forestry authority"),
    "LegalStatistics_Authority": ("central", "Legal statistics body"),
    "CWR": ("agency", "Water Resources Committee"),
    "CLR": ("agency", "Land Resources Committee"),
    "ERC": ("agency", "Eco-control Committee"),
    "CFWL": ("agency", "Forestry and Wildlife Committee"),
    "CRNM": ("agency", "Natural Monopolies Committee"),
    "CCHPU": ("agency", "Utilities Committee"),
    "National_Hydrogeological_Service": ("agency", "Hydrogeological Service"),
    "National_Hydrometeorological_Service": ("agency", "Hydromet Service"),
    "Scientific_Analytical_Organizations": ("agency", "Scientific organisations"),
    "MVRI_Subordinate_Organization": ("agency", "Ministry subordinate"),
    "Accredited_Laboratories": ("agency", "Accredited laboratories"),
    "Accredited_Organization": ("agency", "Technical auditors"),
    "Water_Management_Organizations": ("agency", "Water management orgs"),
    "BWM": ("basin", "Basin Water Management"),
    "BM_Reg": ("basin", "BM_Reg"),
    "Basin_Council": ("basin", "Basin Council"),
    "Akimat_Region": ("basin", "Akimats"),
    "Oblast_Representative_Body": ("basin", "Maslikhats"),
    "Vodokanal_Local": ("basin", "Local water utilities"),
    "Water_Users_Consumers": ("users", "Water users"),
    "Legal_Entity": ("users", "Legal persons"),
    "Public": ("users", "General public"),
    "Enterprise_Agro": ("users", "Farm enterprises"),
    "Irrigation_Condominium": ("users", "Irrigation condominiums"),
    "Fishery_Operators": ("users", "Fishery operators"),
    "Forestry_Entities": ("users", "Forestry entities"),
    "Owner": ("users", "Structure owners"),
    "Industrial_Agro_Polluters": ("users", "Polluters"),
    "Public_and_Enterprises": ("users", "Public and enterprises"),
    "Ecosystems": ("users", "Ecosystems"),
    "Biodiversity": ("users", "Biodiversity"),
    "Environment": ("users", "Environment"),
    "Земли_и_Леса": ("users", "Land and forests"),
}
NOT_ACTORS = {"N/A", "TBD", "Def", "", "None"}

MINISTRY_CODES = ["MVRI", "MA", "MENR", "MNE", "MIC", "ME"]
SOURCES = {"Wat_Code": "Water Code", "Eco_Code": "Ecological Code"}


def clean(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def na(v):
    v = clean(v)
    return "" if v in ("N/A", "None") else v


def split_actors(v):
    return [a for a in (x.strip() for x in re.split(r"[;,]", clean(v))) if a not in NOT_ACTORS]


def deontic(v):
    v = clean(v)
    if v.startswith("Must_Not"):
        return "Must not"
    if v.startswith("Must"):
        return "Must"
    if v.startswith("May"):
        return "May"
    return ""


def article(sid):
    m = re.search(r"Art\.?(\d+)", sid)
    return int(m.group(1)) if m else None


def read_statements(wb):
    rows = list(wb["Raw_IG_Coding"].iter_rows(values_only=True))
    head = [clean(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        sid = clean(d["Statement_ID"])
        if not sid:
            continue
        out.append({
            "id": sid,
            "doc": clean(d["Source_Doc"]),
            "art": article(sid),
            "text": clean(d["Original_Text"]),
            "as": clean(d["Action_Situation"]).upper(),
            "attr": na(d["Attribute [A]"]),
            "deontic": deontic(d["Deontic [D]"]),
            "aim": na(d["Aim [I]"]),
            "object": na(d["Direct_Object [O]"]),
            "context": na(d["Context [C]"]),
            "ruleType": clean(d["Rule_Type"]),
            "ostrom": na(d["Ostrom_Class"]),
            "src": split_actors(d["SNA_Source"]),
            "tgt": split_actors(d["SNA_Target"]),
        })
    return out


def read_codebook(wb):
    book = {}
    for r in wb["ReadMe_Codebook"].iter_rows(values_only=True):
        code, name, role = (clean(x) for x in (list(r) + [None] * 3)[:3])
        if code and name and code != "Actor code":
            book[code] = (name, role)
    return book


def build_actors(statements, book):
    held = collections.Counter()
    targeted = collections.Counter()
    for s in statements:
        held.update(s["src"])
        targeted.update(s["tgt"])
    actors = []
    for code in sorted(set(held) | set(targeted)):
        if code not in ACTORS:
            raise SystemExit(f"Actor code {code!r} has no group/label in ACTORS")
        group, short = ACTORS[code]
        name, role = book.get(code, ("", ""))
        actors.append({
            "code": code, "short": short, "name": name or short, "role": role,
            "group": group, "held": held[code], "targeted": targeted[code],
        })
    return actors


def build_links(statements):
    pairs = collections.Counter()
    for s in statements:
        for a in s["src"]:
            for b in s["tgt"]:
                if a != b:
                    pairs[(a, b)] += 1
    return [{"source": a, "target": b, "n": n} for (a, b), n in sorted(pairs.items())]


def read_ministries():
    xml = zipfile.ZipFile(DOCX).read("word/document.xml").decode("utf8")

    def text(s):
        return clean(re.sub(r"<[^>]+>", "", s))

    captions = re.findall(
        r"Table S\d+: Full listing of organizational units within the (Ministry of .+?) of the Republic",
        text(re.sub(r"</w:p>", "\n", re.sub(r"<w:tbl>.*?</w:tbl>", "", xml, flags=re.S))),
    )
    tables = re.findall(r"<w:tbl>.*?</w:tbl>", xml, flags=re.S)
    if len(captions) != len(tables) or len(tables) != len(MINISTRY_CODES):
        raise SystemExit("Supplementary document layout changed: tables and captions no longer match")
    out = []
    for code, name, tbl in zip(MINISTRY_CODES, captions, tables):
        rows = [[text(c) for c in re.findall(r"<w:tc>.*?</w:tc>", r, flags=re.S)]
                for r in re.findall(r"<w:tr[ >].*?</w:tr>", tbl, flags=re.S)]
        bodies = collections.OrderedDict()
        for _, category, unit, parent in rows[1:]:
            bodies.setdefault(parent, collections.OrderedDict()).setdefault(category, []).append(unit)
        out.append({
            "code": code, "name": name, "units": len(rows) - 1,
            "bodies": [{"name": b, "units": sum(len(u) for u in cats.values()),
                        "categories": [{"name": c, "units": u} for c, u in cats.items()]}
                       for b, cats in bodies.items()],
        })
    return out


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    statements = read_statements(wb)
    book = read_codebook(wb)
    data = {
        "actionSituations": [{"code": c, "name": n, "blurb": b} for c, n, b in ACTION_SITUATIONS],
        "groups": [{"id": g, "name": n} for g, n in GROUPS],
        "sources": SOURCES,
        "statements": statements,
        "actors": build_actors(statements, book),
        "links": build_links(statements),
        "ministries": read_ministries(),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf8")
    print(f"{len(statements)} statements, {len(data['actors'])} actors, "
          f"{len(data['links'])} links, {len(data['ministries'])} ministries "
          f"-> {OUT.relative_to(ROOT)} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
